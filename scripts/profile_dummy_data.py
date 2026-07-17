"""One-off profiling script: full-pass stats over the 3 dummy Excel files.

Used to design the mapping to leads.json / campaigns.json schema.
"""
from __future__ import annotations

import os
from collections import Counter, defaultdict

import openpyxl

FOLDER = r"c:\Github_Projects\Hackathon_AWS\Hackathon-Dig-In\Data Dummy"


def profile_file1():
    print("=" * 80)
    print("FILE 1: Dummy Kiro 1.xlsx")
    wb = openpyxl.load_workbook(os.path.join(FOLDER, "Dummy Kiro 1.xlsx"), read_only=True, data_only=True)
    ws = wb["Sheet1"]

    jenis_leads = Counter()
    media_blasting = Counter()
    by_program = defaultdict(lambda: {"count": 0, "qris_yes": 0, "qris_no": 0, "min_start": None, "max_start": None})

    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        nama_program = row[1]
        jl = row[2]
        mb = row[5]
        periode_start = str(row[8])
        flag_program = row[10]
        qris_first_flag = row[28]  # 'YES'/'NO'

        jenis_leads[jl] += 1
        media_blasting[mb] += 1

        key = (nama_program, flag_program)
        d = by_program[key]
        d["count"] += 1
        if qris_first_flag == "YES":
            d["qris_yes"] += 1
        else:
            d["qris_no"] += 1
        if d["min_start"] is None or periode_start < d["min_start"]:
            d["min_start"] = periode_start
        if d["max_start"] is None or periode_start > d["max_start"]:
            d["max_start"] = periode_start

    print("jenis_leads (full):", jenis_leads)
    print("media_blasting (full):", media_blasting)
    print("by_program:")
    for k, d in by_program.items():
        print(" ", k, d)
    wb.close()


def profile_file2():
    print("=" * 80)
    print("FILE 2: Dummy Kiro 2.xlsx")
    wb = openpyxl.load_workbook(os.path.join(FOLDER, "Dummy Kiro 2.xlsx"), read_only=True, data_only=True)
    ws = wb["Sheet"]

    by_group = defaultdict(lambda: {"count": 0, "takeup": 0, "min_start": None, "max_start": None})
    sub_crs = Counter()
    segment_div_owner = Counter()

    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        campaign_name = row[2]
        jenis_leads = row[3]
        periode_start = str(row[6])
        segment_div_owner_v = row[7]
        sub_crs_v = row[8]
        norek_tapenas = row[11]

        sub_crs[sub_crs_v] += 1
        segment_div_owner[segment_div_owner_v] += 1

        key = (campaign_name, jenis_leads)
        d = by_group[key]
        d["count"] += 1
        if norek_tapenas != "NULL":
            d["takeup"] += 1
        if d["min_start"] is None or periode_start < d["min_start"]:
            d["min_start"] = periode_start
        if d["max_start"] is None or periode_start > d["max_start"]:
            d["max_start"] = periode_start

    print("by_group:")
    for k, d in by_group.items():
        print(" ", k, d)
    print("sub_crs:", sub_crs)
    print("segment_div_owner:", segment_div_owner)
    wb.close()


def profile_file3():
    print("=" * 80)
    print("FILE 3: Dummy Kiro 3.xlsx")
    wb = openpyxl.load_workbook(os.path.join(FOLDER, "Dummy Kiro 3.xlsx"), read_only=True, data_only=True)
    ws = wb["Sheet"]

    status = Counter()
    flag_balrun = Counter()
    lifegoals_filled = 0
    total = 0
    unique_cifs = set()

    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        total += 1
        status[row[5]] += 1
        flag_balrun[row[7]] += 1
        cif = row[6]
        unique_cifs.add(cif)
        lifegoals_account = row[11]
        if lifegoals_account not in (None, "NULL", ""):
            lifegoals_filled += 1

    print("total rows:", total)
    print("unique cifs:", len(unique_cifs))
    print("status:", status)
    print("flag_balrun:", flag_balrun)
    print("lifegoals_account filled:", lifegoals_filled)
    wb.close()


if __name__ == "__main__":
    profile_file1()
    profile_file2()
    profile_file3()
