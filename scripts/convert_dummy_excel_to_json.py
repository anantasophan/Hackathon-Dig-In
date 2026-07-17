"""Convert raw dummy Excel files into JSON (one JSON array per file).

This is a one-off utility script for the hackathon data-prep workflow.
It does NOT map columns to the CampaignInsightGenerator ``leads.json``
schema -- it preserves the original Excel columns as-is so we can
inspect the data in JSON form first and decide the mapping afterwards.

Usage:
    python scripts/convert_dummy_excel_to_json.py
"""

from __future__ import annotations

import datetime
import json
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SRC_DIR = ROOT / "Data Dummy"
OUT_DIR = ROOT / "Data Dummy" / "json"

FILES = [
    ("Dummy Kiro 1.xlsx", "dummy_kiro_1.json"),
    ("Dummy Kiro 2.xlsx", "dummy_kiro_2.json"),
    ("Dummy Kiro 3.xlsx", "dummy_kiro_3.json"),
]


def _json_default(value: object) -> str:
    """Serialize datetime/date objects to ISO strings for json.dump."""
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.isoformat()
    return str(value)


def convert_file(src_path: Path, out_path: Path) -> int:
    """Convert a single Excel file to a JSON array of row objects.

    Args:
        src_path: Path to the source .xlsx file.
        out_path: Path to write the resulting .json file.

    Returns:
        Number of data rows written.
    """
    wb = openpyxl.load_workbook(src_path, read_only=True, data_only=True)
    sheet_name = wb.sheetnames[0]
    ws = wb[sheet_name]

    row_iter = ws.iter_rows(values_only=True)
    headers = next(row_iter)
    headers = [str(h) if h is not None else f"col_{i}" for i, h in enumerate(headers)]

    count = 0
    with out_path.open("w", encoding="utf-8") as fh:
        fh.write("[\n")
        first = True
        for row in row_iter:
            if row[0] is None and all(v is None for v in row):
                continue
            record = dict(zip(headers, row))
            if not first:
                fh.write(",\n")
            json.dump(record, fh, ensure_ascii=False, default=_json_default)
            first = False
            count += 1
        fh.write("\n]\n")

    wb.close()
    return count


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for src_name, out_name in FILES:
        src_path = SRC_DIR / src_name
        out_path = OUT_DIR / out_name
        print(f"Converting {src_name} -> json/{out_name} ...")
        n = convert_file(src_path, out_path)
        size_mb = out_path.stat().st_size / (1024 * 1024)
        print(f"  {n} rows written, {size_mb:.1f} MB")


if __name__ == "__main__":
    main()
